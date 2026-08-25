import os
import re
import openpyxl
from pathlib import Path

polling_dir = Path("C:/Users/techb/Desktop/electoral-portal-main/poling addres")
excel_files = sorted(polling_dir.glob("*.xlsx"))

master_map = {}

for f in excel_files:
    wb = openpyxl.load_workbook(f, data_only=True)
    for sheetname in wb.sheetnames:
        ws = wb[sheetname]
        for row in ws.iter_rows(values_only=True):
            if not row or not any(row): continue
            r_str = [str(c).strip() if c is not None else '' for c in row]
            
            # Find part number and station name
            part_no = None
            for idx in range(len(r_str)):
                cell_val = r_str[idx]
                if re.match(r'^\d{1,3}$', cell_val):
                    val = int(cell_val)
                    if 1 <= val <= 200:
                        part_no = str(val)
                        break
                        
            if not part_no:
                continue
                
            # Station Name
            station = ""
            area = ""
            district = ""
            taluk = ""
            
            for cell_val in r_str:
                if 'college' in cell_val.lower() or 'school' in cell_val.lower() or 'office' in cell_val.lower() or 'kacheri' in cell_val.lower() or 'hall' in cell_val.lower() or 'court' in cell_val.lower() or 'panchayat' in cell_val.lower() or 'room' in cell_val.lower():
                    if len(cell_val) > len(station):
                        station = cell_val
                elif 'hobli' in cell_val.lower() or 'ward' in cell_val.lower() or 'entire' in cell_val.lower() or 'taluk' in cell_val.lower() or 'village' in cell_val.lower():
                    if len(cell_val) > len(area):
                        area = cell_val
                        
            station = re.sub(r'\s+', ' ', station.replace('\n', ' ')).strip()
            area = re.sub(r'\s+', ' ', area.replace('\n', ' ')).strip()
            
            if station and not station.isdigit():
                master_map[part_no] = {
                    'station': station,
                    'area': area or f"Part {part_no} Constituency Area, Karnataka"
                }

print("====================================================")
print(f"BUILT COMPLETE POLLING MAP: {len(master_map)} PARTS LOADED")
print("====================================================")
for p in sorted(master_map.keys(), key=int):
    print(f"Part {p:>3}: Station: '{master_map[p]['station'][:50]}' | Area: '{master_map[p]['area'][:50]}'")

cur_close = None
