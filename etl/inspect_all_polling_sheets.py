import openpyxl
from pathlib import Path

p = Path("C:/Users/techb/Desktop/electoral-portal-main/poling addres")
for f in sorted(p.glob("*.xlsx")):
    print(f"\n====================================================")
    print("FILE:", f.name)
    wb = openpyxl.load_workbook(f, data_only=True)
    for sheetname in wb.sheetnames:
        print(" SHEET:", sheetname)
        ws = wb[sheetname]
        count = 0
        for r in list(ws.iter_rows(values_only=True))[:10]:
            print("   ", [str(c)[:40] if c is not None else '' for c in r])
            count += 1
