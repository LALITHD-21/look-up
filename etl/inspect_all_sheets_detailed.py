import openpyxl
from pathlib import Path

p = Path("C:/Users/techb/Desktop/electoral-portal-main/poling addres")

for f in sorted(p.glob("*.xlsx")):
    wb = openpyxl.load_workbook(f, data_only=True)
    print(f"\n====================================================")
    print("FILE:", f.name, "| Sheets:", wb.sheetnames)
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        parts = []
        for r in ws.iter_rows(values_only=True):
            if not r: continue
            for cell in r:
                if cell is not None and str(cell).strip().isdigit():
                    v = int(str(cell).strip())
                    if 1 <= v <= 250:
                        parts.append(v)
        if parts:
            print(f"  - Sheet '{sheet}': Parts range {min(parts)} to {max(parts)} (Total part numbers found: {len(set(parts))})")

print("====================================================")
