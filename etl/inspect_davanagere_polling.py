import openpyxl

wb = openpyxl.load_workbook("C:/Users/techb/Desktop/electoral-portal-main/poling addres/Davanagere- Sheet2.xlsx", data_only=True)
for sheet in wb.sheetnames:
    ws = wb[sheet]
    for row in ws.iter_rows(values_only=True):
        if not row or not any(row): continue
        r_str = [str(c).strip() if c is not None else '' for c in row]
        if any(r_str):
            print(f"[{sheet}]", r_str[:8])
