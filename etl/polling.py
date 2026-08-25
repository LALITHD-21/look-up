"""
ETL Module: Polling Station Address Loader & Mapping
Loads master polling station mapping from poling addres/ directory.
"""

import os
import re
import logging
from pathlib import Path
import openpyxl

logger = logging.getLogger(__name__)

FALLBACK_POLLING_MAP = {
    '92': {
        'polling_station_name': 'Kalidasa Composite Pre- University College, Tumkur',
        'polling_address': 'Entire Ward No.01 to 05, Tumkur Taluk, Tumkur District, Karnataka'
    },
    '93': {
        'polling_station_name': 'Government Model Higher Primary School, Shishuvihara Compound, Tumkur',
        'polling_address': 'Entire Ward No.06 to 10, Tumkur Taluk, Tumkur District, Karnataka'
    },
    '94': {
        'polling_station_name': 'Siddaganga Pre University College,B H Road Gandhinagara, Tumkur',
        'polling_address': 'Entire Ward No.11 to 15, Tumkur Taluk, Tumkur District, Karnataka'
    },
    '95': {
        'polling_station_name': 'Government Higher Primary School, Shanthi nagara ASK Palya',
        'polling_address': 'Entire Ward No.16 to 20, Tumkur Taluk, Tumkur District, Karnataka'
    },
    '96': {
        'polling_station_name': 'Sri Siddaganga Kannada Elementory Higher Primary School, Room No-2, Tumkur',
        'polling_address': 'Entire Ward No.21 to 25, Tumkur Taluk, Tumkur District, Karnataka'
    },
    '97': {
        'polling_station_name': 'Nalanda convent and high school, sapthagiri extension, Room No-1, Tumkur',
        'polling_address': 'Entire Ward No.26 to 30, Tumkur Taluk, Tumkur District, Karnataka'
    },
    '98': {
        'polling_station_name': 'Government Model Higher Primary School, Kyathsandra, Room No-1, Tumkur.',
        'polling_address': 'Entire Ward No.31 to 35, Tumkur Taluk, Tumkur District, Karnataka'
    },
    '99': {
        'polling_station_name': 'Court Hall, Taluk Office Tumkur',
        'polling_address': 'Entire Kasaba Hobli - Rural, Tumkur Taluk, Tumkur District, Karnataka'
    },
    '100': {
        'polling_station_name': 'Govt. Higher Primary School, Gulur',
        'polling_address': 'Entire Gulur Hobli, Tumkur Taluk, Tumkur District, Karnataka'
    },
    '101': {
        'polling_station_name': 'Ganapathi High School, Hebbur',
        'polling_address': 'Entire Hebbur Hobli, Tumkur Taluk, Tumkur District, Karnataka'
    },
    '102': {
        'polling_station_name': 'Govt. Higher Primary School, Urdigere',
        'polling_address': 'Entire Urdigere Hobli, Tumkur Taluk, Tumkur District, Karnataka'
    },
    '103': {
        'polling_station_name': 'Kuvempu Govt. Primary School, Bellavi',
        'polling_address': 'Entire Bellavi Hobli, Tumkur Taluk, Tumkur District, Karnataka'
    },
    '104': {
        'polling_station_name': 'Govt. Model Higher Primary School, Kora',
        'polling_address': 'Entire Kora Hobli, Tumkur Taluk, Tumkur District, Karnataka'
    },
}


def load_polling_parts_mapping(polling_dir: str = '../poling addres') -> dict:
    """
    Read all Excel mapping files from polling_dir.
    Returns dict mapping part_number (str) -> {polling_station_name, polling_address}.
    """
    polling_path = Path(polling_dir)
    master_map = dict(FALLBACK_POLLING_MAP)

    if not polling_path.exists():
        logger.warning(f"Polling directory not found: {polling_dir}")
        return master_map

    excel_files = sorted(polling_path.glob('*.xlsx'))
    logger.info(f"Loading polling address mapping from {len(excel_files)} files in {polling_dir}...")

    for f in excel_files:
        try:
            wb = openpyxl.load_workbook(f, data_only=True)
            for s in wb.sheetnames:
                ws = wb[s]
                for row in ws.iter_rows(values_only=True):
                    if not row or not any(row):
                        continue
                    r_str = [str(c).strip() if c is not None else '' for c in row]

                    part_no = None
                    for idx in [3, 2, 0, 1]:
                        if idx < len(r_str) and re.match(r'^\d{1,3}$', r_str[idx]):
                            val = int(r_str[idx])
                            if 1 <= val <= 250:
                                if idx == 3 or not part_no:
                                    part_no = str(val)

                    if not part_no:
                        continue

                    district = r_str[1] if len(r_str) > 1 else ''
                    taluk = r_str[2] if len(r_str) > 2 else ''
                    station = r_str[4] if len(r_str) > 4 else (r_str[3] if len(r_str) > 3 else '')
                    area = r_str[5] if len(r_str) > 5 else (r_str[4] if len(r_str) > 4 else '')

                    station = re.sub(r'\s+', ' ', station.replace('\n', ' ')).strip()
                    area = re.sub(r'\s+', ' ', area.replace('\n', ' ')).strip()
                    area = re.sub(r'(\d{1,2})\d{3,5}$', r'\1', area).strip()
                    district = re.sub(r'\s+', ' ', district).strip()
                    taluk = re.sub(r'\s+', ' ', taluk).strip()

                    if not station or 'part name' in station.lower() or 'sl. no' in station.lower() or station.isdigit():
                        continue

                    full_address = area
                    if taluk and taluk.lower() not in area.lower():
                        full_address += f', {taluk} Taluk'
                    if district and district.lower() not in area.lower():
                        full_address += f', {district} District'
                    full_address += ', Karnataka'

                    master_map[part_no] = {
                        'polling_station_name': station,
                        'polling_address': full_address
                    }
        except Exception as e:
            logger.error(f"Failed to read polling map file {f.name}: {e}")

    logger.info(f"✓ Master polling map ready with {len(master_map)} parts.")
    return master_map
